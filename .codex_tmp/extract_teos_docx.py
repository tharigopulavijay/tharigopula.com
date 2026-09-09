from pathlib import Path
from docx import Document

src = Path(r"C:\Users\vijay\Downloads\TEOS_Master_Prompt_by_Prompt_Build_Playbook.docx")
out = Path(r"E:\my businesses\tharigopula\.codex_tmp\TEOS_playbook_extracted.txt")
doc = Document(src)

lines = []
for i, p in enumerate(doc.paragraphs, 1):
    text = p.text.strip()
    if text:
        lines.append(f"[P{i:04d}] [{p.style.name}] {text}")

for ti, table in enumerate(doc.tables, 1):
    lines.append(f"\n[TABLE {ti}]")
    for ri, row in enumerate(table.rows, 1):
        cells = [" ".join(c.text.split()) for c in row.cells]
        lines.append(f"R{ri:03d}: " + " || ".join(cells))

out.parent.mkdir(parents=True, exist_ok=True)
out.write_text("\n".join(lines), encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
print(f"extracted_lines={len(lines)} output={out}")
